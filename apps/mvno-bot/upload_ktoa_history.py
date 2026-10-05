"""
upload_ktoa_history.py
KOTA.xlsx → Firestore ktoa_daily 일괄 업로드 스크립트

[수정 이력]
v1.0 | 2026-03-23 | 최초 작성
  - 통계 시트: 일별 MVNO in/out, MNO out, 순증감 업로드
  - 영업일수 실 data 시트: bw_manual(추정기존), bw_performance(실적기반) 업로드
  - 누적값은 엑셀 원본값 직접 사용
"""

import sys
sys.path.insert(0, '/home/mclee_cecilia/.local/lib/python3.12/site-packages')

import openpyxl
from datetime import datetime, timedelta
from google.cloud import firestore

# ── Firestore 초기화 (Cloud Shell gcloud auth 사용) ─────────────────────────
db = firestore.Client(project='mvno-484509', database='mvno-data')

# ── 엑셀 로드 ────────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook('KOTA.xlsx', data_only=True)

# ============================================================
# STEP 1: 영업일수 데이터 추출 (영업일수 실 data 시트)
# ============================================================
ws_bw = wb['영업일수 실 data']
rows_bw = list(ws_bw.iter_rows(min_row=1, max_row=ws_bw.max_row, values_only=True))

bw_map = {}  # {날짜str: {'bw_manual': x, 'bw_performance': y}}

i = 0
while i < len(rows_bw):
    row = rows_bw[i]
    # 월 헤더 행 찾기
    if row[0] and isinstance(row[0], str) and '월' in str(row[0]):
        date_row  = rows_bw[i+1] if i+1 < len(rows_bw) else None
        perf_row  = rows_bw[i+3] if i+3 < len(rows_bw) else None  # 영업일수 실적기반
        est_row   = rows_bw[i+4] if i+4 < len(rows_bw) else None  # 영업일수 추정기존

        if date_row and perf_row and est_row:
            for j, date_val in enumerate(date_row):
                if isinstance(date_val, datetime):
                    date_str  = date_val.strftime('%Y-%m-%d')
                    perf_bw   = perf_row[j]
                    est_bw    = est_row[j]
                    # 둘 중 하나라도 있으면 저장
                    if isinstance(perf_bw, (int, float)) or isinstance(est_bw, (int, float)):
                        entry = {}
                        if isinstance(perf_bw, (int, float)):
                            entry['bw_performance'] = round(float(perf_bw), 4)
                        if isinstance(est_bw, (int, float)):
                            entry['bw_manual'] = round(float(est_bw), 4)
                        bw_map[date_str] = entry
    i += 1

print(f"영업일수 추출: {len(bw_map)}개 날짜")

# ============================================================
# STEP 2: 통계 데이터 추출 및 Firestore 업로드
# ============================================================
ws_stat = wb['통계']

# 컬럼 매핑 (0-indexed)
# col0: 일자
# col1~3: MVNO in 일별 (SM, KM, LM)
# col4~6: MVNO in 누적 (SM, KM, LM)
# col7~9: MNO Out 일별 (S, K, L)
# col10~12: MNO Out 누적 (S, K, L)
# col13~15: 순증감 일별 (S, K, L)
# col16~18: 순증감 누적 (S, K, L)
# col19~21: MVNO Out 일별 (S, K, L)
# col22~24: MVNO Out 누적 (S, K, L)
# col25: 영업일수 추정기존 (통계 시트에도 있음)

uploaded = 0
skipped  = 0

for row in ws_stat.iter_rows(min_row=3, max_row=ws_stat.max_row, values_only=True):
    date_val = row[0]
    if not date_val:
        continue

    # 날짜 파싱
    if isinstance(date_val, datetime):
        date_str = date_val.strftime('%Y-%m-%d')
    elif isinstance(date_val, str):
        try:
            date_str = datetime.strptime(date_val, '%Y-%m-%d').strftime('%Y-%m-%d')
        except:
            continue
    else:
        continue

    # 미래 날짜 스킵
    if date_str > '2026-03-21':
        skipped += 1
        continue

    # SM 값이 없으면 스킵 (휴일 등)
    sm_val = row[1]
    if sm_val is None:
        skipped += 1
        continue

    def safe_int(v): return int(v) if isinstance(v, (int, float)) else 0

    # 당일 데이터
    mvno_in   = {'SM': safe_int(row[1]),  'KM': safe_int(row[2]),  'LM': safe_int(row[3]),
                 '계': safe_int(row[1]) + safe_int(row[2]) + safe_int(row[3])}
    cum_mvno_in = {'SM': safe_int(row[4]), 'KM': safe_int(row[5]),  'LM': safe_int(row[6]),
                   '계': safe_int(row[4]) + safe_int(row[5]) + safe_int(row[6])}
    mno_out   = {'S':  safe_int(row[7]),  'K':  safe_int(row[8]),  'L':  safe_int(row[9]),
                 '계': safe_int(row[7]) + safe_int(row[8]) + safe_int(row[9])}
    cum_mno_out = {'S': safe_int(row[10]), 'K':  safe_int(row[11]), 'L':  safe_int(row[12]),
                   '계': safe_int(row[10]) + safe_int(row[11]) + safe_int(row[12])}
    net_change  = {'SM': safe_int(row[13]), 'KM': safe_int(row[14]), 'LM': safe_int(row[15]),
                   '계': safe_int(row[13]) + safe_int(row[14]) + safe_int(row[15])}
    cum_net     = {'SM': safe_int(row[16]), 'KM': safe_int(row[17]), 'LM': safe_int(row[18]),
                   '계': safe_int(row[16]) + safe_int(row[17]) + safe_int(row[18])}
    mvno_out    = {'SM': safe_int(row[19]), 'KM': safe_int(row[20]), 'LM': safe_int(row[21]),
                   '계': safe_int(row[19]) + safe_int(row[20]) + safe_int(row[21])}
    cum_mvno_out = {'SM': safe_int(row[22]), 'KM': safe_int(row[23]), 'LM': safe_int(row[24]),
                    '계': safe_int(row[22]) + safe_int(row[23]) + safe_int(row[24])}

    # 영업일수
    bw_info = bw_map.get(date_str, {})
    bw_manual      = bw_info.get('bw_manual', None)
    bw_performance = bw_info.get('bw_performance', None)
    # 통계 시트 col25도 확인
    if bw_manual is None and len(row) > 25 and row[25] is not None:
        bw_manual = round(float(row[25]), 4)

    # total 계산
    total_mvno = mvno_in['계']
    total_mno  = mno_out['계']  # MNO→MVNO
    total      = total_mvno + total_mno

    payload = {
        'date':           date_str,
        'reference_time': '20시 00분 이전',
        'source':         'excel_import',
        # 당일
        'mvno_in':        mvno_in,
        'mvno_out':       mvno_out,
        'mno_out':        mno_out,
        'net_change':     net_change,
        # 누적
        'cum_mvno_in':    cum_mvno_in,
        'cum_mvno_out':   cum_mvno_out,
        'cum_mno_out':    cum_mno_out,
        'cum_net':        cum_net,
        # 영업일수
        'bw_manual':      bw_manual,
        'bw_performance': bw_performance,
        # raw
        'total':          total,
        'total_mno':      total_mno,
        'total_mvno':     total_mvno,
    }

    # None 제거
    payload = {k: v for k, v in payload.items() if v is not None}

    db.collection('ktoa_daily').document(date_str).set(payload, merge=True)
    uploaded += 1

    if uploaded % 50 == 0:
        print(f"  진행중... {uploaded}건 업로드")

print(f"\n✅ 완료! 업로드: {uploaded}건, 스킵: {skipped}건")