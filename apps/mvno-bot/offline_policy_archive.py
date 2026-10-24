"""
오프라인 유통(부성/비전) 정책 원문 아카이브 - 실적·사유 분석용 [추가 20261024]

cafe_collector.py의 텔레그램 메시지(7G/11G 두 구간)는 그대로 두고, 같은 게시글을
"분석용"으로 전부 저장한다 (사용자 결정 2026-10-24: 메시지는 그대로, 수집은 전부).

  - 원본 첨부파일(엑셀·이미지)을 GCS에 보관: gs://mvno-484509-offline-policy/{source}/{article_id}/
  - 요금제 표의 모든 행·칸 (구분, 요금제명, 기본료, 할인, 월청구, 음성·부가·문자·데이터·속도,
    정책 4종 = 내국인/외국인 × 010신규/MNP, 비고 등)
      vision(비전): 첨부 엑셀(.xlsx)을 직접 파싱 (2024-01~ 모든 글에 있음, LLM 불필요)
      busung(부성): 사업자별 이미지 → Gemini 전사 (엑셀 없음)
  - 본문: 차수·적용 시작일·사업자별 변경 내용·차감 조건·기간 한정 특별정책
  - 시트/이미지 안 안내문(환수 기준, 유심비·가입비, 차감 등)

Firestore: offline_policy_rounds/{source}_{article_id}  (글마다 1문서, 덮어쓰지 않음 - 수정글도 별도 보존)
  적용 종료일(effective_to)은 다음 차수 시작일 전날로 recompute_effective_periods()가 채운다.

실행:
  새 글 - cafe_collector.process_new_article()가 archive_article()을 호출 (실패해도 메시지 흐름엔 영향 없음)
  소급 - python offline_policy_archive.py backfill {busung|vision} --since 2025-01-01
         (Cloud Run 병렬 task: CLOUD_RUN_TASK_INDEX/COUNT로 글을 나눠 처리, 이미 저장된 글은 건너뜀)
"""
import io
import os
import re
import sys
import json
import time
import urllib.parse
from datetime import datetime, timedelta, timezone

import requests

KST = timezone(timedelta(hours=9))
UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36')
BUCKET = 'mvno-484509-offline-policy'
COLLECTION = 'offline_policy_rounds'
PARSE_VERSION = '1'

# 표에서 쓰는 사업자 표기 → 메시지 라벨 (cafe_collector의 row_order 라벨과 동일 체계)
CARRIER_KEYWORDS = [
    ('조이텔', ('조이텔', 'JOYTEL')),
    ('KTM', ('KTM', 'KT M', 'KT엠', '케이티엠', 'M mobile')),
    ('스카이', ('스카이', 'SKYLIFE', 'SKY LIFE')),
    ('텔링크', ('세븐', 'SK7', 'SK 7', '7MOBILE', '7 MOBILE', '텔링크')),
    ('프리티', ('프리티', '프리텔')),
    ('유모비', ('유모비', 'U+알뜰', 'U+ 알뜰', 'U MOBILE', 'UMOBILE', '유플러스알뜰', 'U+유모바일')),
    ('헬로', ('헬로', 'HELLO')),
    ('스노우', ('스노우',)),
    ('이야기', ('이야기',)),
]
# 라벨 → KTOA 사업자(ktoa_mvno_brand_in.brand / operator_registry.name) - 실적과 붙일 때 쓰는 키
CARRIER_KTOA = {
    'KTM': 'KT엠모바일', '스카이': 'KT스카이라이프', '헬로': 'LG헬로비전',
    '텔링크': 'SK텔링크(재판매)', '유모비': '미디어로그', '조이텔': '조이텔', '프리티': '프리텔레콤',
}
# 분석 대상이 아닌 시트 (원본 파일은 GCS에 그대로 있음)
SKIP_SHEET = re.compile(r'요약|단말|선불|유선|인터넷|와이파이|상조|TV', re.I)


def _carrier_of(text):
    t = (text or '').upper().replace(' ', '')
    for label, keys in CARRIER_KEYWORDS:
        if any(k.upper().replace(' ', '') in t for k in keys):
            return label
    return None


def _carriers_in(text):
    """텍스트에 나온 사업자들을 등장 순서대로 (시트명 '후불 유심 (세븐,프리텔레콤SKT,KT,U+)' → [텔링크, 프리티])"""
    t = (text or '').upper().replace(' ', '')
    found = []
    for label, keys in CARRIER_KEYWORDS:
        pos = [t.find(k.upper().replace(' ', '')) for k in keys if k.upper().replace(' ', '') in t]
        if pos:
            found.append((min(pos), label))
    return [lb for _, lb in sorted(found)]


def _num(v):
    """'20,000원' / 20000 / 49060.00000000001 → int, 숫자가 아니면 None"""
    if isinstance(v, (int, float)):
        return int(round(v))
    m = re.fullmatch(r'\s*([\d,]+)\s*원?\s*', str(v or ''))
    return int(m.group(1).replace(',', '')) if m else None


def data_gb_month(text):
    """'11GB+일2GB' → 71 (매일×30), '7GB + 5GB' → 12, '500MB' → 0.49, '일 5GB' → 150, 무제한 → 9999"""
    t = str(text or '')
    if '무제한' in t:
        return 9999.0
    gb = 0.0
    for m in re.finditer(r'(일|매일)?\s*([\d.]+)\s*(GB|G|MB)(?![a-z])', t, re.I):
        v = float(m.group(2)) / (1024 if m.group(3).upper() == 'MB' else 1)
        gb += v * (30 if m.group(1) else 1)
    return round(gb, 2) if gb else None


# ── 첨부파일 수집 ───────────────────────────────────────────────────────

def fetch_article_full(source, article_id, driver, cookies):
    """게시글 본문 + 첨부파일(다운로드 링크) 목록. driver는 쿠키가 이미 들어간 셀레니움 드라이버."""
    cafe_id, menu_id = source['cafe_id'], source['menu_id']
    driver.get(f'https://cafe.naver.com/ca-fe/cafes/{cafe_id}/articles/{article_id}'
               f'?menuid={menu_id}&referrerAllArticles=false&fromNext=true')
    time.sleep(4)
    body_text = driver.execute_script('return document.body.innerText')
    files = driver.execute_script(
        'return [...document.querySelectorAll("a.se-file-save-button")].map(a => a.href)')
    if len(body_text) < 400:
        raise RuntimeError(f'본문을 못 읽음({len(body_text)}자) - 네이버 세션 쿠키 만료 가능성')
    return {'body_text': body_text, 'file_urls': files}


def download_attachment(url, cookies):
    r = requests.get(url, cookies=cookies, headers={'User-Agent': UA, 'Referer': 'https://cafe.naver.com/'},
                     timeout=90)
    r.raise_for_status()
    cd = r.headers.get('content-disposition', '')
    name = urllib.parse.unquote(cd.split("filename*=UTF-8''")[-1]) if "filename*=" in cd else os.path.basename(url)
    return name.strip('"'), r.content


def _gcs_bucket():
    from google.cloud import storage
    return storage.Client(project='mvno-484509').bucket(BUCKET)


# ── 엑셀 파싱 (vision) ───────────────────────────────────────────────────

def _sheet_grid(ws):
    """셀값 dict {(r,c): str}, 병합셀은 범위 전체에 같은 값"""
    g = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ''):
                g[(c.row, c.column)] = c.value
    for mr in ws.merged_cells.ranges:
        v = ws.cell(mr.min_row, mr.min_col).value
        if v in (None, ''):
            continue
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                g[(r, c)] = v
    return g


def _clean(v):
    return re.sub(r'\s+', ' ', str(v)).strip()


_HDR_NAME = re.compile(r'^(요금제명?|《?Best ?요금제》?|★?BEST상품★?\[?주력요금제\]?)$', re.I)


def _is_header_row(g, r, max_c):
    vals = [_clean(g[(r, c)]) for c in range(1, max_c + 1) if (r, c) in g]
    return any(_HDR_NAME.match(v.replace(' ', '')) for v in vals) and \
        any('기본료' in v or '청구' in v or '음성' in v for v in vals)


def _col_role(label):
    """머리글 → 표준 필드명 (없으면 None = extra로 보존)"""
    L = label.replace(' ', '')
    last = label.split(' / ')[-1].replace(' ', '')
    if 'MNP' in L.upper() or '번호이동' in L or '010' in L or '신규' in last:
        if any(k in L for k in ('결합', 'DATA', '데이터')):
            return None
        side = ('cn' if '중국인' in L else
                'for' if ('외국인정책' in L or L.startswith('외국인')) and '내국인' not in L and '내/외' not in L else 'dom')
        kind = 'mnp' if ('MNP' in L.upper() or '번호이동' in L) else 'new'
        return f'{side}_{kind}'
    if any(_HDR_NAME.match(x.replace(' ', '')) for x in label.split(' / ')):
        return 'name'           # '요금제명 / 일반 유심'처럼 둘째 줄에 부제가 붙는 형식 포함
    if last == '구분':
        return 'category'
    if '청구' in L or '(A-B)' in L or '평생기본료' in L:
        return 'monthly_fee'
    if last.startswith('기본료'):
        return 'base_fee'
    if '할인' in last:
        return 'discount'
    if last.startswith('음성'):
        return 'voice'
    if last.startswith('부가통화') or last.startswith('영상'):
        return 'extra_voice'
    if last.startswith('문자'):
        return 'sms'
    if last.startswith('데이터') and '프로모션' not in last:
        return 'data'
    if '안심옵션' in L or 'QOS' in L.upper() or '속도' in last:
        return 'speed'
    if last in ('비고', '기타', '필수사항'):
        return 'note'
    return None


def parse_policy_sheet(ws, default_carrier):
    """시트 하나 → (plans, notes). 한 시트에 표가 여러 개(사업자 여러 곳)여도 표마다 사업자를 따로 잡음."""
    g = _sheet_grid(ws)
    if not g:
        return [], []
    max_r = max(r for r, _ in g)
    max_c = min(max(c for _, c in g), 80)
    used_rows = set()
    plans = []
    sheet_carriers = _carriers_in(ws.title)
    table_no, last_table_end = 0, 0

    r = 1
    while r <= max_r:
        if not _is_header_row(g, r, max_c):
            r += 1
            continue
        # 머리글 묶음: 연속된 머리글 행 + 바로 위 그룹 머리글(내국인/외국인 정책)
        band = [r]
        while band[-1] + 1 <= max_r and _is_header_row(g, band[-1] + 1, max_c):
            band.append(band[-1] + 1)
        nxt = band[-1] + 1
        # 둘째 줄 머리글(음성/문자/010/MNP)이 '요금제명' 없이 오는 형식
        if any(_clean(g.get((nxt, c), '')) in ('음성', '010 (신규)', 'MNP (번호이동)') for c in range(1, max_c + 1)):
            band.append(nxt)
        top = band[0]
        # 사업자: 표 바로 위 안내문/제목(직전 표가 끝난 뒤 행만) → 없으면 시트 이름에 나온 사업자를 표 순서대로
        # (예: '후불 유심 (U+알뜰,헬로)' 첫 표=유모비, 둘째 표=헬로)
        carrier = None
        for rr in range(top - 1, max(top - 8, last_table_end, 0), -1):
            texts = ' '.join(_clean(g[(rr, c)]) for c in range(1, max_c + 1) if (rr, c) in g)
            carrier = _carrier_of(texts)
            if carrier:
                break
        if not carrier:
            carrier = sheet_carriers[min(table_no, len(sheet_carriers) - 1)] if sheet_carriers else default_carrier
        table_no += 1

        labels, roles = {}, {}
        for c in range(1, max_c + 1):
            parts = []
            for rr in band:
                v = _clean(g.get((rr, c), ''))
                if v and v not in parts:
                    parts.append(v)
            if parts:
                labels[c] = ' / '.join(parts)
        seen_role, seen_label = set(), set()
        for c in sorted(labels):
            role = _col_role(labels[c])
            if role and role not in seen_role:
                roles[c] = role
                seen_role.add(role)
            elif labels[c] not in seen_label:      # 역할 없거나 중복(예: 요약 열) → 원문 머리글로 보존
                roles[c] = 'x:' + labels[c][:40]
            seen_label.add(labels[c])
        if 'name' not in seen_role:
            r = band[-1] + 1
            continue
        used_rows.update(band)

        rr = band[-1] + 1
        blank, last_data_row = 0, band[-1]
        while rr <= max_r and not _is_header_row(g, rr, max_c):
            row = {}
            for c, role in roles.items():
                if (rr, c) in g:
                    row[role] = g[(rr, c)]
            name = _clean(row.get('name', ''))
            pol = [row.get(k) for k in ('dom_new', 'dom_mnp', 'for_new', 'for_mnp', 'cn_new', 'cn_mnp')]
            if not name:
                blank += 1
                if blank >= 2:
                    break
                rr += 1
                continue
            blank = 0
            if not any(_num(v) is not None for v in pol) and _num(row.get('monthly_fee')) is None:
                rr += 1         # 표 사이 안내문 줄
                continue
            used_rows.add(rr)
            last_data_row = rr
            plan = {'carrier': carrier, 'name': name}
            for k in ('category', 'voice', 'extra_voice', 'sms', 'data', 'speed', 'note'):
                if k in row:
                    plan[k] = _clean(row[k])
            for k in ('base_fee', 'discount', 'monthly_fee', 'dom_new', 'dom_mnp', 'for_new', 'for_mnp',
                      'cn_new', 'cn_mnp'):
                if k in row:
                    n = _num(row[k])
                    plan[k] = n
                    if n is None and _clean(row[k]) not in ('', '-'):
                        plan[k + '_text'] = _clean(row[k])     # 예: '내국인 개통불가'
            extra = {k[2:]: _clean(v) for k, v in row.items() if k.startswith('x:') and _clean(v) not in ('', '-')}
            if extra:
                plan['extra'] = extra
            plan['data_gb_month'] = data_gb_month(plan.get('data'))
            plans.append(plan)
            rr += 1
        last_table_end = last_data_row
        r = rr

    # 표 밖 안내문 (환수 기준, 유심비·가입비, 차감 등) - 병합셀 중복 제거, 행 순서대로
    notes, seen = [], set()
    for r in range(1, max_r + 1):
        if r in used_rows:
            continue
        for c in range(1, max_c + 1):
            v = g.get((r, c))
            if not isinstance(v, str):
                continue
            t = _clean(v)
            if len(t) >= 4 and t not in seen:
                seen.add(t)
                notes.append(t)
    return plans, notes[:120]


def parse_policy_xlsx(content):
    import openpyxl
    import warnings
    warnings.filterwarnings('ignore')
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    plans, sheet_notes, sheets = [], {}, []
    for ws in wb.worksheets:
        sheets.append(ws.title)
        if SKIP_SHEET.search(ws.title):
            continue
        p, n = parse_policy_sheet(ws, _carrier_of(ws.title) or ws.title)
        for x in p:
            x['sheet'] = ws.title
        plans += p
        if n:
            sheet_notes[ws.title] = n
    return plans, sheet_notes, sheets


# ── 이미지 전사 (busung, 엑셀 없는 vision 글) ───────────────────────────────

# 열 순서 고정 배열로 받아 출력 토큰을 줄임 (JSON 키 반복 시 KTM처럼 45행 표는 응답이 길어 504 발생)
IMAGE_COLS = ['category', 'name', 'base_fee', 'discount', 'monthly_fee', 'voice', 'sms', 'extra_voice', 'data',
              'speed', 'dom_new', 'dom_mnp', 'for_new', 'for_mnp', 'extra', 'note']
IMAGE_PROMPT = (
    '이 이미지는 알뜰폰 사업자의 유통 정책표입니다. 표에 보이는 값을 판단 없이 그대로 옮겨 적으세요.\n'
    'JSON 객체 하나로만 응답: {"carrier": 표 상단 사업자 로고/제목, "period": 상단 적용기간 문구, '
    '"rows": [[...], ...], "notes": [...]}\n'
    'rows = 요금제 행마다 배열 하나, 순서: [구분(좌측 묶음 칸), 요금제명, 기본료, 할인, 월청구금액(프로모션가·표에 금액이 하나면 여기), '
    '음성, 문자, 부가통화(영상외), 데이터 원문(예 "11GB+일2GB"), QOS(안심옵션 속도), 내국인010신규, 내국인MNP, 외국인010신규, 외국인MNP, '
    '그밖의칸("머리글:값; 머리글:값", 예 아무나결합 추가데이터), 비고(필수사항 칸)]\n'
    '금액은 숫자만(127000). 빈칸·"-"는 null. "내국인 개통불가" 같은 글자는 그대로. '
    '내국인/외국인 블록 혼동 금지(왼쪽=내국인, 오른쪽=외국인). 병합된 칸은 해당 행마다 값을 채울 것.\n'
    'notes = 표 밖 안내문 줄마다 하나씩 원문 그대로(차감 조건, 유심비·가입비, 환수 기준 등).'
)


def transcribe_policy_image(image_bytes, mime='image/png'):
    """Gemini REST 직접 호출 (google-generativeai 0.3.2는 타임아웃·thinking 설정이 안 돼 긴 표에서 504)"""
    import base64
    key = os.environ.get('GEMINI_API_KEY') or os.environ.get('GOOGLE_API_KEY', '')
    url = 'https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent'
    body = {
        'contents': [{'parts': [{'text': IMAGE_PROMPT},
                                {'inline_data': {'mime_type': mime, 'data': base64.b64encode(image_bytes).decode()}}]}],
        'generationConfig': {'temperature': 0, 'maxOutputTokens': 32768, 'responseMimeType': 'application/json',
                             'thinkingConfig': {'thinkingBudget': 0}},
    }
    last = None
    for attempt in range(4):
        try:
            r = requests.post(url, params={'key': key}, json=body, timeout=300)
            if r.status_code in (429, 500, 503, 504):
                raise RuntimeError(f'HTTP {r.status_code}')
            r.raise_for_status()
            text = r.json()['candidates'][0]['content']['parts'][0]['text']
            obj = json.loads(re.sub(r'^```(?:json)?\s*|\s*```$', '', text.strip()))
            plans = []
            for row in obj.get('rows') or []:
                if isinstance(row, list):
                    d = dict(zip(IMAGE_COLS, row + [None] * (len(IMAGE_COLS) - len(row))))
                    if isinstance(d.get('extra'), str) and d['extra']:
                        d['extra'] = dict(kv.split(':', 1) for kv in d['extra'].split(';') if ':' in kv)
                    plans.append(d)
            obj['plans'] = plans
            return obj
        except Exception as e:
            last = e
            time.sleep(10 * (attempt + 1))
    raise last


def _normalize_image_plans(obj, default_carrier):
    carrier = _carrier_of(obj.get('carrier')) or default_carrier
    out = []
    for p in obj.get('plans') or []:
        if not p.get('name'):
            continue
        plan = {'carrier': carrier, 'name': _clean(p['name'])}
        for k in ('category', 'voice', 'extra_voice', 'sms', 'data', 'speed', 'note'):
            if p.get(k) not in (None, ''):
                plan[k] = _clean(p[k])
        for k in ('base_fee', 'discount', 'monthly_fee', 'dom_new', 'dom_mnp', 'for_new', 'for_mnp'):
            if p.get(k) not in (None, ''):
                n = _num(p[k])
                plan[k] = n
                if n is None:
                    plan[k + '_text'] = _clean(p[k])
        if isinstance(p.get('extra'), dict) and p['extra']:
            plan['extra'] = {_clean(k)[:40]: _clean(v) for k, v in p['extra'].items() if v not in (None, '', '-')}
        plan['data_gb_month'] = data_gb_month(plan.get('data'))
        out.append(plan)
    return out


# ── 본문 파싱 ─────────────────────────────────────────────────────────────

def clean_body(body_text, subject=None):
    """카페 화면 전체 텍스트 → 게시글 본문만 (작성자·조회수 줄 뒤 ~ 첨부파일 목록/게시판 목록 앞)"""
    t = body_text
    m = re.search(r'조회\s*[\d,]+\s*\n\s*댓글\s*\d+\s*URL 복사\s*\n', t)
    if m:
        t = t[m.end():]
    t = re.sub(r'^\s*첨부파일 모아보기 \d+\s*\n', '', t)
    for stop in ('\n첨부파일\n', '님의 게시글 더보기'):
        k = t.find(stop)
        if k > 0:
            t = t[:k]
    lines = [ln.strip() for ln in t.replace('\u200b', '').splitlines()]
    return '\n'.join(ln for ln in lines if ln)


def parse_effective_from(subject, body, write_date):
    """적용 시작일 (게시일 기준으로 연·월 보정). 제목 '10월 24일' / '(24일~)' / '24일부터' / 본문 '2026년 10월 24일 ~'"""
    y, mo = write_date.year, write_date.month
    for text in (subject, body[:600]):
        m = re.search(r'(?:(\d{4})년\s*)?(\d{1,2})월\s*(\d{1,2})일', text or '')
        if m and not re.search(r'\d+월\s*\d+차', (text or '')[m.start():m.end() + 1]):
            yy = int(m.group(1)) if m.group(1) else y
            mm, dd = int(m.group(2)), int(m.group(3))
            if not m.group(1) and mm == 12 and mo == 1:
                yy -= 1
            elif not m.group(1) and mm == 1 and mo == 12:
                yy += 1
            try:
                return datetime(yy, mm, dd).strftime('%Y-%m-%d')
            except ValueError:
                pass
        m = re.search(r'\(?\s*(\d{1,2})일\s*(?:부터|~)', text or '')
        if m:
            dd = int(m.group(1))
            mm_m = re.search(r'(\d{1,2})월', text or '')
            mm = int(mm_m.group(1)) if mm_m else mo
            yy = y + (1 if mm == 1 and mo == 12 else -1 if mm == 12 and mo == 1 else 0)
            try:
                return datetime(yy, mm, dd).strftime('%Y-%m-%d')
            except ValueError:
                pass
    return write_date.strftime('%Y-%m-%d')


_AMT = re.compile(r'(-?\s*\d+(?:\.\d+)?)\s*만\s*원?')
_PERIOD = re.compile(r'(\d{1,2})/(\d{1,2})\s*~\s*(?:(\d{1,2})/)?(\d{1,2})|(\d{1,2})일\s*까지')


def parse_body_notices(body):
    """사업자별 공지 줄 + 조건(차감/특별정책/면제) 추출. 형식 두 가지:
       부성 'N. 사업자 (10/24~ )' + '☞ ...', 비전 '◆ 사업자' + '- ...' / ': ...' / '※ ...'"""
    notices, cur = {}, None
    for ln in body.splitlines():
        s = ln.strip()
        m = re.match(r'^(?:\d+\.\s*|◆\s*)([^☞\-:※]+?)(?:\s*\(.*\))?\s*$', s)
        if m and len(s) <= 40 and _carrier_of(m.group(1)):
            cur = _carrier_of(m.group(1))
            continue
        if s.startswith('◈') and re.search(r'\d+\s*차', s):
            continue                         # '◈케이엘모바일_부성정책_9차◈' 제목줄
        if s.startswith('◈') and not _carrier_of(s):
            cur = '공통'                     # '◈ 전 통신사 전월대비 단가변동'
            notices.setdefault(cur, []).append(s.lstrip('◈ ').strip())
            continue
        if s.startswith(('―', '<', '[', '감사합니다')) or re.match(r'^[─―━-]{3,}', s):
            if s.startswith(('―', '<', '감사합니다')):
                cur = None
            continue
        if s.startswith('※') and _carrier_of(s) and _carrier_of(s) != cur:
            notices.setdefault(_carrier_of(s), []).append(s.lstrip('※ ').strip())   # 다른 사업자 섹션 안의 ※ 줄
        elif re.match(r'^(☞|-|:|※|\*|·|•|★)', s) and cur:
            notices.setdefault(cur, []).append(s.lstrip('☞-:*·•★ ').strip())

    conditions = []
    for label, lines in notices.items():
        for s in lines:
            kind = '차감' if '차감' in s else '특별정책' if re.search(r'특별|프로모션|최대|지급중', s) else \
                   '면제' if '면제' in s else '정책변경' if re.search(r'상승|인상|하락|인하|추가|변경|변동|대폭|축소|확대|출시|판가|환수', s) else None
            if not kind:
                continue
            c = {'carrier': label, 'kind': kind, 'text': s}
            amts = [float(a.replace(' ', '')) for a in _AMT.findall(s)]
            if amts:
                c['amount_won'] = int(abs(max(amts, key=abs)) * 10000)
            pm = _PERIOD.search(s)
            if pm:
                c['period'] = pm.group(0)
            conditions.append(c)
    return notices, conditions


# ── 저장 ──────────────────────────────────────────────────────────────────

def _db():
    from google.cloud import firestore
    return firestore.Client(project='mvno-484509', database='mvno-data')


def _extract_round_no(subject):
    m = re.search(r'(\d+(?:-\d+)?)\s*차', subject or '')
    return m.group(1) if m else ''


def archive_article(source_key, source, article, driver, cookies, body_text=None, force=False):
    """게시글 1건을 원본 보관 + 파싱해서 offline_policy_rounds에 저장. 반환: 저장한 레코드 요약"""
    db = _db()
    article_id = int(article['article_id'])
    doc_id = f'{source_key}_{article_id}'
    ref = db.collection(COLLECTION).document(doc_id)
    if not force and ref.get(field_paths=['parse']).exists:
        return {'doc_id': doc_id, 'skipped': True}

    detail = fetch_article_full(source, article_id, driver, cookies)
    body_raw = body_text or detail['body_text']
    write_date = datetime.fromtimestamp(article['write_ts_ms'] / 1000, tz=KST)
    bucket = _gcs_bucket()

    files, xlsx_list, images, errors = [], [], [], []
    for url in detail['file_urls']:
        try:
            name, content = download_attachment(url, cookies)
        except Exception as e:
            errors.append(f'download: {e}')
            continue
        path = f'{source_key}/{article_id}/{name}'
        bucket.blob(path).upload_from_string(content)
        ext = os.path.splitext(name)[1].lower()
        files.append({'name': name, 'gcs': f'gs://{BUCKET}/{path}', 'bytes': len(content)})
        if ext in ('.xlsx', '.xlsm'):
            xlsx_list.append((name, content))
        elif ext in ('.png', '.jpg', '.jpeg'):
            images.append((name, content))

    plans, sheet_notes, sheets, method = [], {}, [], None
    for name, content in xlsx_list:
        try:
            p, n, s = parse_policy_xlsx(content)
            plans += p
            sheet_notes.update(n)
            sheets += s
            method = 'xlsx'
        except Exception as e:
            errors.append(f'xlsx {name}: {e}')
    if not plans:
        # 엑셀이 없거나 실패 → 이미지 전사 (요약본·QR·안내 이미지는 제외)
        # 이미지마다 Gemini 1회 (6장 순차면 수 분 걸려 병렬 처리)
        from concurrent.futures import ThreadPoolExecutor
        filename_map = source.get('filename_to_label', {})
        targets = []
        for name, content in images:
            stem = os.path.splitext(name)[0]
            if re.search(r'요약|QR|선택동의|캡처|안내', stem, re.I):
                continue
            default = _carrier_of(stem) or filename_map.get(stem) or stem
            mime = 'image/jpeg' if name.lower().endswith(('.jpg', '.jpeg')) else 'image/png'
            targets.append((name, content, default, mime))

        def _one(t):
            name, content, default, mime = t
            try:
                return name, default, transcribe_policy_image(content, mime), None
            except Exception as e:
                return name, default, None, e

        with ThreadPoolExecutor(max_workers=6) as ex:
            results = list(ex.map(_one, targets))
        for name, default, obj, err in results:
            if err is not None or not isinstance(obj, dict):
                errors.append(f'image {name}: {err or "JSON 아님"}')
                continue
            p = _normalize_image_plans(obj, default)
            for x in p:
                x['sheet'] = name
            plans += p
            if obj.get('notes'):
                sheet_notes[name] = [str(t) for t in obj['notes']][:80]
            if obj.get('period'):
                sheet_notes.setdefault(name, []).insert(0, f"[적용기간] {obj['period']}")
            method = 'gemini'

    body = clean_body(body_raw, article['subject'])
    notices, conditions = parse_body_notices(body)
    for p in plans:
        p['carrier_ktoa'] = CARRIER_KTOA.get(p['carrier'])

    record = {
        'source': source_key,
        'article_id': article_id,
        'subject': article['subject'],
        'writer': article.get('writer', ''),
        'write_date': write_date,
        'round_no': _extract_round_no(article['subject']),
        'is_revision': bool(re.search(r'수정', article['subject'])),
        'effective_from': parse_effective_from(article['subject'], body, write_date),
        'effective_to': None,
        'body': body[:20000],
        'notices': notices,
        'conditions': conditions,
        'files': files,
        'sheets': sheets,
        'sheet_notes': sheet_notes,
        'plans': plans,
        'carriers': sorted({p['carrier'] for p in plans}),
        'parse': {'method': method, 'plan_count': len(plans), 'errors': errors[:20], 'version': PARSE_VERSION},
        'created_at': datetime.now(KST),
    }
    ref.set(record)
    return {'doc_id': doc_id, 'plans': len(plans), 'method': method, 'errors': errors}


def recompute_effective_periods(source_key):
    """effective_to = 다음 글의 적용 시작일 전날. 같은 날 여러 글(수정글)이면 앞 글은 superseded=True."""
    db = _db()
    docs = list(db.collection(COLLECTION).where('source', '==', source_key)
                .select(['effective_from', 'article_id', 'effective_to', 'superseded']).stream())
    rows = sorted(((d.to_dict(), d.reference) for d in docs),
                  key=lambda x: (x[0].get('effective_from') or '', x[0].get('article_id') or 0))
    batch, n = db.batch(), 0
    for i, (d, ref) in enumerate(rows):
        nxt = rows[i + 1][0] if i + 1 < len(rows) else None
        superseded = bool(nxt and nxt.get('effective_from') == d.get('effective_from'))
        if nxt and not superseded:
            to = (datetime.strptime(nxt['effective_from'], '%Y-%m-%d') - timedelta(days=1)).strftime('%Y-%m-%d')
        elif superseded:
            to = d.get('effective_from')
        else:
            to = None
        if d.get('effective_to') != to or bool(d.get('superseded')) != superseded:
            batch.update(ref, {'effective_to': to, 'superseded': superseded})
            n += 1
            if n % 400 == 0:
                batch.commit()
                batch = db.batch()
    batch.commit()
    return n


# ── 소급 ──────────────────────────────────────────────────────────────────

def list_articles(source, since):
    cafe_id, menu_id = source['cafe_id'], source['menu_id']
    out = []
    for page in range(1, 40):
        r = requests.get(
            f'https://apis.naver.com/cafe-web/cafe-boardlist-api/v1/cafes/{cafe_id}/menus/{menu_id}/articles',
            params={'page': page, 'pageSize': 50, 'sortBy': 'TIME', 'viewType': 'L'},
            headers={'User-Agent': UA, 'Referer': f'https://cafe.naver.com/f-e/cafes/{cafe_id}/menus/{menu_id}'},
            timeout=15)
        r.raise_for_status()
        items = r.json()['result']['articleList']
        if not items:
            break
        for it in items:
            a = it['item']
            if datetime.fromtimestamp(a['writeDateTimestamp'] / 1000, tz=KST).strftime('%Y-%m-%d') < since:
                return out
            out.append({'article_id': a['articleId'], 'subject': a['subject'],
                        'writer': a['writerInfo']['nickName'], 'write_ts_ms': a['writeDateTimestamp']})
    return out


def open_driver():
    import cafe_collector as cc
    aut, ses = cc._get_cookies()
    d = cc._get_chrome_driver()
    d.get('https://cafe.naver.com/robots.txt')
    d.add_cookie({'name': 'NID_AUT', 'value': aut, 'domain': '.naver.com', 'path': '/'})
    d.add_cookie({'name': 'NID_SES', 'value': ses, 'domain': '.naver.com', 'path': '/'})
    return d, {'NID_AUT': aut, 'NID_SES': ses}


def backfill(source_key, since, until=None, force=False, ids=None):
    import cafe_collector as cc
    source = cc.SOURCES[source_key]
    arts = list_articles(source, since)
    if until:
        arts = [a for a in arts if datetime.fromtimestamp(a['write_ts_ms'] / 1000, tz=KST).strftime('%Y-%m-%d') <= until]
    if ids:
        arts = [a for a in arts if a['article_id'] in ids]
    arts.sort(key=lambda a: a['article_id'])
    idx, cnt = int(os.environ.get('CLOUD_RUN_TASK_INDEX', 0)), int(os.environ.get('CLOUD_RUN_TASK_COUNT', 1))
    mine = arts[idx::cnt]
    print(f'[backfill] {source_key} since={since} 전체 {len(arts)}건 중 task {idx}/{cnt} 담당 {len(mine)}건', flush=True)
    driver, cookies = open_driver()
    ok = fail = 0
    try:
        for i, a in enumerate(mine):
            try:
                res = archive_article(source_key, source, a, driver, cookies, force=force)
                ok += 1
                print(f'[backfill] {i+1}/{len(mine)} #{a["article_id"]} {a["subject"][:30]} → {res}', flush=True)
            except Exception as e:
                fail += 1
                print(f'[backfill] FAIL #{a["article_id"]} {a["subject"][:30]}: {e}', flush=True)
                if 'chrome' in str(e).lower() or 'session' in str(e).lower():
                    try:
                        driver.quit()
                    except Exception:
                        pass
                    driver, cookies = open_driver()
    finally:
        driver.quit()
    print(f'[backfill] 완료 ok={ok} fail={fail}', flush=True)
    if idx == 0 or cnt == 1:
        print(f'[backfill] effective_to 갱신 {recompute_effective_periods(source_key)}건', flush=True)


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('cmd', choices=['backfill', 'periods'])
    ap.add_argument('source', choices=['busung', 'vision'])
    ap.add_argument('--since', default='2025-01-01')
    ap.add_argument('--until')
    ap.add_argument('--ids', help='쉼표로 구분한 article_id만')
    ap.add_argument('--force', action='store_true')
    a = ap.parse_args()
    if a.cmd == 'periods':
        print(recompute_effective_periods(a.source))
    else:
        backfill(a.source, a.since, a.until, a.force,
                 [int(x) for x in a.ids.split(',')] if a.ids else None)
