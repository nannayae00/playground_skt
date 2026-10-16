"""
verify_prepaid_vs_excel.py  v1.0
작성일: 2026-10-16

팀 엑셀('선불정보 DB구축 통계' .xlsm)의 월 시트와 수집값을 날짜·사업자·지표 단위로 대조한다.
엑셀 월 시트는 사이트 표와 같은 구조라 ktoa_prepaid_parser.parse_grid()로 그대로 읽는다.

사용법:
    # Firestore(ktoa_prepaid_daily)와 비교
    python3 verify_prepaid_vs_excel.py 선불정보.xlsm
    python3 verify_prepaid_vs_excel.py 선불정보.xlsm --sheets 26년10월,26년09월
    # 수집기 DRY_RUN 결과(JSON)와 비교
    python3 verify_prepaid_vs_excel.py 선불정보.xlsm --json /tmp/ktoa_prepaid_out
"""

import os
import sys
import glob
import json
import argparse
from collections import defaultdict

import openpyxl

from ktoa_prepaid_parser import parse_grid, METRICS

SKIP_SHEETS = {"제어판", "종합"}


def load_excel(path: str, sheets: list[str] | None) -> tuple[dict, dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = sheets or [s for s in wb.sheetnames if s not in SKIP_SHEETS]
    days, names_by_key = {}, {}
    for name in names:
        ws = wb[name]
        grid = [list(r) for r in ws.iter_rows(min_row=1, max_row=40, values_only=True)]
        p = parse_grid(grid)
        days.update(p["days"])
        names_by_key.update({b["key"]: b["name"] for b in p["blocks"]})
    return days, names_by_key


def load_json(dir_or_file: str) -> dict:
    files = [dir_or_file] if os.path.isfile(dir_or_file) else glob.glob(os.path.join(dir_or_file, "*.json"))
    days = {}
    for f in sorted(files):
        with open(f, encoding="utf-8") as fh:
            days.update(json.load(fh)["days"])
    return days


def compare(excel: dict, collected: dict, names: dict) -> int:
    dates = sorted(set(excel) & set(collected))
    only_excel = sorted(set(excel) - set(collected))
    mismatches = []
    missing_blocks = defaultdict(list)   # 엑셀엔 있고 수집값엔 없는 블록
    extra_blocks = defaultdict(list)     # 수집값에만 있는 블록 (엑셀에 아직 안 만든 신규 사업자 등)

    for d in dates:
        e, c = excel[d], collected[d]
        for k in e.keys() - c.keys():
            if any(e[k].values()):
                missing_blocks[k].append(d)
        for k in c.keys() - e.keys():
            if any(c[k].values()):
                extra_blocks[k].append(d)
        for k in e.keys() & c.keys():
            for m in METRICS:
                if e[k][m] != c[k][m]:
                    mismatches.append((d, k, m, e[k][m], c[k][m]))

    print(f"비교 일자 {len(dates)}일 ({dates[0] if dates else '-'} ~ {dates[-1] if dates else '-'})")
    if only_excel:
        print(f"수집값에 없는 엑셀 일자 {len(only_excel)}일: {only_excel[:5]}{' ...' if len(only_excel) > 5 else ''}")
    print(f"값 불일치 {len(mismatches)}건")
    for d, k, m, ev, cv in mismatches[:50]:
        print(f"  {d} {names.get(k, k)} {m}: 엑셀 {ev:,} / 수집 {cv:,} (차 {cv - ev:+,})")
    for title, blocks in (("엑셀에만 있는 블록(실적 있음)", missing_blocks),
                          ("수집값에만 있는 블록(실적 있음)", extra_blocks)):
        if blocks:
            print(f"{title} {len(blocks)}개")
            for k, ds in list(blocks.items())[:20]:
                print(f"  {names.get(k, k)}: {len(ds)}일 (첫 {ds[0]})")
    ok = not mismatches and not missing_blocks
    print("결과: 일치" if ok else "결과: 차이 있음")
    return 0 if ok else 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsm")
    ap.add_argument("--sheets", help="쉼표로 구분한 월 시트 이름 (기본: 전체 월 시트)")
    ap.add_argument("--json", help="DRY_RUN JSON 폴더 또는 파일 (없으면 Firestore)")
    a = ap.parse_args()

    excel, names = load_excel(a.xlsm, a.sheets.split(",") if a.sheets else None)
    if not excel:
        print("엑셀에서 일자 데이터를 찾지 못함")
        return 1
    if a.json:
        collected = load_json(a.json)
    else:
        from ktoa_prepaid_firestore import load_prepaid_days
        collected = load_prepaid_days(min(excel), max(excel))
    return compare(excel, collected, names)


if __name__ == "__main__":
    sys.exit(main())
